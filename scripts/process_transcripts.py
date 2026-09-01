#!/usr/bin/env python3
"""
Read transcript JSON from the Excel workbook, fill participant_text_clean (user-only text),
and session_number_per_employee (1, 2, 3, … per employee_id by time).

Sessions with no transcription (empty / null) are skipped: no participant text, and they do
not receive a session number (they are not counted in the per-employee sequence). If such a
row wrongly has a session number from older runs, it is cleared on the next write.

By default only updates rows that need participant text and/or session cells (including
clears and corrections). Use --all-rows to recompute and write both columns for every data row.

Close Microsoft Excel before running so the file is not locked.

The data sheet is resolved automatically: prefers ``Database Sessions``, otherwise a sheet
whose name contains "transcript", or set ``TRANSCRIPTS_SHEET_NAME``. Saves via a temp file
then replace to reduce risk of a half-written workbook on errors.

Usage:
  python process_transcripts.py
  python process_transcripts.py --input "Master-Supabase Transcripts.xlsx"
  python process_transcripts.py --all-rows
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

import workbook_utils

HEADER_ROW = 2  # 1-based: row with session_id, employee_id, …
DATA_START_ROW = HEADER_ROW + 1


def _participant_text_from_transcription(raw: object) -> str:
    if raw is None or (isinstance(raw, float) and pd.isna(raw)):
        return ""
    s = str(raw).strip()
    if not s or s.lower() == "null":
        return ""
    try:
        data = json.loads(s)
    except json.JSONDecodeError:
        return ""
    messages = data.get("messages")
    if not isinstance(messages, list):
        return ""
    parts: list[str] = []
    for m in messages:
        if not isinstance(m, dict):
            continue
        if str(m.get("role", "")).lower() != "user":
            continue
        text = m.get("text")
        if text is None:
            continue
        chunk = str(text).strip()
        if chunk:
            parts.append(chunk)
    return "\n\n".join(parts)


def _session_sort_timestamp(df: pd.DataFrame) -> pd.Series:
    started = pd.to_datetime(df["started_at"], utc=True, errors="coerce")
    ended = pd.to_datetime(df["ended_at"], utc=True, errors="coerce")
    ts = started.where(started.notna(), ended)
    return ts


def _is_missing_text(val: object) -> bool:
    if val is None or (isinstance(val, float) and pd.isna(val)):
        return True
    s = str(val).strip()
    if not s or s.lower() in ("nan", "none", "null"):
        return True
    return False


def _transcription_present(raw: object) -> bool:
    if raw is None or (isinstance(raw, float) and pd.isna(raw)):
        return False
    s = str(raw).strip()
    return bool(s) and s.lower() not in ("null", "none", "nan")


def _computed_session_numbers(df: pd.DataFrame, has_transcript: pd.Series) -> pd.Series:
    """1, 2, 3, … per employee_id by time, only for rows with a present transcription."""
    out = pd.Series(pd.NA, index=df.index, dtype="Int64")
    if not has_transcript.any():
        return out
    sub = df.loc[has_transcript].copy()
    sub["_ts"] = _session_sort_timestamp(sub)
    sub = sub.sort_values(
        by=["employee_id", "_ts", "session_id"],
        kind="mergesort",
        na_position="last",
    )
    sub["_n"] = sub.groupby("employee_id", sort=False).cumcount() + 1
    out.loc[sub.index] = sub["_n"]
    return out


def _is_missing_session(val: object) -> bool:
    if val is None or (isinstance(val, float) and pd.isna(val)):
        return True
    if isinstance(val, str):
        s = val.strip()
        if not s or s.lower() in ("nan", "none", "null"):
            return True
        try:
            return int(float(s)) < 1
        except ValueError:
            return True
    try:
        return int(val) < 1  # type: ignore[arg-type]
    except (TypeError, ValueError):
        return True


def main() -> int:
    parser = argparse.ArgumentParser(description="Fill participant_text_clean and session numbers.")
    parser.add_argument(
        "--input",
        type=Path,
        default=workbook_utils.DEFAULT_TRANSCRIPTS,
        help="Path to Master-Supabase Transcripts.xlsx",
    )
    parser.add_argument(
        "--all-rows",
        action="store_true",
        help="Recompute and write both columns for every data row (full refresh).",
    )
    args = parser.parse_args()
    path = args.input.expanduser().resolve()
    if not path.is_file():
        print(f"File not found: {path}", file=sys.stderr)
        return 1

    xl = pd.ExcelFile(path)
    sheet_name = workbook_utils.transcripts_sheet_name(list(xl.sheet_names))
    df = pd.read_excel(xl, sheet_name=sheet_name, header=HEADER_ROW - 1)
    if "transcription" not in df.columns or "employee_id" not in df.columns:
        print("Expected columns 'transcription' and 'employee_id' not found.", file=sys.stderr)
        return 1
    for col in ("participant_text_clean", "session_number_per_employee"):
        if col not in df.columns:
            print(f"Expected column {col!r} not found.", file=sys.stderr)
            return 1

    has_transcript = df["transcription"].map(_transcription_present)
    computed_text = df["transcription"].map(_participant_text_from_transcription)
    computed_session = _computed_session_numbers(df, has_transcript)

    if args.all_rows:
        needs_text = pd.Series(True, index=df.index)
        needs_session_write = pd.Series(True, index=df.index)
    else:
        # Only fill text when there is transcript and extractable user content.
        has_extractable_user_text = computed_text.str.len() > 0
        needs_text = (
            df["participant_text_clean"].map(_is_missing_text)
            & has_transcript
            & has_extractable_user_text
        )
        cur_sess = df["session_number_per_employee"]
        cur_num = pd.to_numeric(cur_sess, errors="coerce")
        comp_num = computed_session.astype("float")
        # Rows without transcript: clear session cell if anything was stored.
        w_clear = ~has_transcript & ~cur_sess.map(_is_missing_session)
        # Rows with transcript: write if missing or differs from recomputed sequence.
        w_fill = has_transcript & (
            cur_sess.map(_is_missing_session)
            | (computed_session.notna() & (cur_num != comp_num))
        )
        needs_session_write = w_clear | w_fill

    wb = load_workbook(path)
    if sheet_name not in wb.sheetnames:
        print(f"Sheet {sheet_name!r} not in workbook. Available: {wb.sheetnames}", file=sys.stderr)
        return 1
    ws = wb[sheet_name]

    header = [ws.cell(HEADER_ROW, c).value for c in range(1, ws.max_column + 1)]
    try:
        col_text = header.index("participant_text_clean") + 1
        col_sess = header.index("session_number_per_employee") + 1
    except ValueError as e:
        print(f"Could not locate output columns in row {HEADER_ROW}: {e}", file=sys.stderr)
        return 1

    n = len(df)
    n_text = int(needs_text.sum())
    n_sess = int(needs_session_write.sum())
    if n_text == 0 and n_sess == 0:
        print(f"No updates needed in {path.name} ({n} rows). Nothing to write.")
        return 0

    for i in range(n):
        excel_row = DATA_START_ROW + i
        if needs_text.iloc[i]:
            t = computed_text.iloc[i]
            ws.cell(excel_row, col_text).value = t if t else None
        if needs_session_write.iloc[i]:
            val = computed_session.iloc[i]
            ws.cell(excel_row, col_sess).value = int(val) if pd.notna(val) else None

    workbook_utils.save_workbook_atomic(wb, path)
    filled_text = int((needs_text & (computed_text.str.len() > 0)).sum())
    print(
        f"Updated {path.name}: wrote participant_text_clean for {n_text} row(s) "
        f"({filled_text} non-empty), session_number_per_employee for {n_sess} row(s). "
        f"Total data rows: {n}."
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
