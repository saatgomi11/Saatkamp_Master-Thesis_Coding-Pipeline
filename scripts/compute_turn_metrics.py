#!/usr/bin/env python3
"""
Fill participant turn metrics from the full interview JSON (transcription).

- participant_turn_count: number of user (participant) messages in messages[].
- avg_words_per_turn: mean word count across those user messages (0 if no user turns).
- longest_response_words: max word count of a single user message.

Skips rows with no usable transcription. By default only fills empty cells; use --all-rows
to overwrite.

Close Excel before running.

Usage:
  python compute_turn_metrics.py
  python compute_turn_metrics.py --all-rows
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

import workbook_utils

HEADER_ROW = 2
DATA_START_ROW = HEADER_ROW + 1

COL_TRANSCRIPTION = "transcription"
COL_TURNS = "participant_turn_count"
COL_AVG = "avg_words_per_turn"
COL_LONGEST = "longest_response_words"


def _transcription_present(raw: object) -> bool:
    if raw is None or (isinstance(raw, float) and pd.isna(raw)):
        return False
    s = str(raw).strip()
    return bool(s) and s.lower() not in ("null", "none", "nan")


def _count_words(text: str) -> int:
    if not text or not str(text).strip():
        return 0
    return len(str(text).split())


def _is_missing(val: object) -> bool:
    if val is None or (isinstance(val, float) and pd.isna(val)):
        return True
    s = str(val).strip()
    return not s or s.lower() in ("nan", "none", "null")


def _turn_stats_from_transcription(raw: object) -> tuple[int, float, int] | None:
    """
    Returns (participant_turn_count, avg_words_per_turn, longest_response_words), or None
    if transcription cannot be parsed or has no messages list.
    """
    if not _transcription_present(raw):
        return None
    s = str(raw).strip()
    try:
        data = json.loads(s)
    except json.JSONDecodeError:
        return None
    msgs = data.get("messages")
    if not isinstance(msgs, list):
        return None
    user_word_counts: list[int] = []
    for m in msgs:
        if not isinstance(m, dict):
            continue
        if str(m.get("role", "")).lower() != "user":
            continue
        t = m.get("text")
        if t is None:
            continue
        user_word_counts.append(_count_words(str(t)))
    n = len(user_word_counts)
    if n == 0:
        return 0, 0.0, 0
    total = sum(user_word_counts)
    avg = total / n
    longest = max(user_word_counts)
    return n, avg, longest


def main() -> int:
    parser = argparse.ArgumentParser(description="Fill participant turn metrics from transcription JSON.")
    parser.add_argument(
        "--input",
        type=Path,
        default=workbook_utils.DEFAULT_TRANSCRIPTS,
        help="Path to Master-Supabase Transcripts.xlsx",
    )
    parser.add_argument(
        "--all-rows",
        action="store_true",
        help="Overwrite the three columns whenever transcription is present.",
    )
    args = parser.parse_args()
    path = args.input.expanduser().resolve()
    if not path.is_file():
        print(f"File not found: {path}", file=sys.stderr)
        return 1

    xl = pd.ExcelFile(path)
    sheet_name = workbook_utils.transcripts_sheet_name(list(xl.sheet_names))
    df = pd.read_excel(xl, sheet_name=sheet_name, header=HEADER_ROW - 1)

    for c in (COL_TRANSCRIPTION, COL_TURNS, COL_AVG, COL_LONGEST):
        if c not in df.columns:
            print(f"Missing column {c!r}.", file=sys.stderr)
            return 1

    wb = load_workbook(path)
    if sheet_name not in wb.sheetnames:
        print(f"Sheet {sheet_name!r} not found.", file=sys.stderr)
        return 1
    ws = wb[sheet_name]
    header = [ws.cell(HEADER_ROW, c).value for c in range(1, ws.max_column + 1)]
    try:
        idx_turns = header.index(COL_TURNS) + 1
        idx_avg = header.index(COL_AVG) + 1
        idx_long = header.index(COL_LONGEST) + 1
    except ValueError as e:
        print(f"Header lookup failed: {e}", file=sys.stderr)
        return 1

    n = len(df)
    updates = 0

    for i in range(n):
        raw = df[COL_TRANSCRIPTION].iloc[i]
        stats = _turn_stats_from_transcription(raw)
        if stats is None:
            continue
        turn_count, avg_words, longest = stats

        excel_row = DATA_START_ROW + i

        if args.all_rows or _is_missing(df[COL_TURNS].iloc[i]):
            ws.cell(excel_row, idx_turns).value = int(turn_count)
            updates += 1
        if args.all_rows or _is_missing(df[COL_AVG].iloc[i]):
            ws.cell(excel_row, idx_avg).value = round(float(avg_words), 2)
            updates += 1
        if args.all_rows or _is_missing(df[COL_LONGEST].iloc[i]):
            ws.cell(excel_row, idx_long).value = int(longest)
            updates += 1

    if updates == 0:
        print(f"No turn-metric updates needed in {path.name} ({n} rows).")
        return 0

    workbook_utils.save_workbook_atomic(wb, path)
    print(f"Updated {path.name}: wrote {updates} cell(s) on sheet {sheet_name!r}.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
