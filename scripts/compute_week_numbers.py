#!/usr/bin/env python3
"""
Fill two week-related columns on the transcript sheet:

week_number (employee session index)
  How many interviews this employee has completed so far, per employee_id, ordered by
  session time (started_at, else ended_at). First qualifying session = 1, second = 2, etc.
  Across all audits. Only rows with a present transcription count. Rows with missing
  employee_id, missing audit_id, or no transcription are skipped (no number).

audit_week (audit calendar week)
  Per audit_id: find the Monday UTC of the calendar week of the earliest session in that audit.
  That week is audit week 1; each later calendar week is 2, 3, … for sessions in that audit.
  Rows with missing audit_id are skipped.

Only writes empty/invalid cells unless --all-rows is passed.

Close Excel before running.

Usage:
  python compute_week_numbers.py
  python compute_week_numbers.py --all-rows
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

import workbook_utils

HEADER_ROW = 2
DATA_START_ROW = HEADER_ROW + 1

COL_WEEK_NUMBER = "week_number"
COL_AUDIT_WEEK = "audit_week"
COL_TRANSCRIPTION = "transcription"


def _session_ts_utc(row: pd.Series) -> pd.Timestamp:
    started = pd.to_datetime(row.get("started_at"), utc=True, errors="coerce")
    ended = pd.to_datetime(row.get("ended_at"), utc=True, errors="coerce")
    t = started if pd.notna(started) else ended
    return t


def _monday_week_start_utc(ts: pd.Timestamp) -> pd.Timestamp:
    if pd.isna(ts):
        return pd.NaT
    day = ts.floor("D")
    return day - pd.Timedelta(days=int(day.dayofweek))


def _id_present(raw: object) -> bool:
    if raw is None or (isinstance(raw, float) and pd.isna(raw)):
        return False
    s = str(raw).strip()
    return bool(s) and s.lower() not in ("null", "none", "nan")


def _transcription_present(raw: object) -> bool:
    if raw is None or (isinstance(raw, float) and pd.isna(raw)):
        return False
    s = str(raw).strip()
    return bool(s) and s.lower() not in ("null", "none", "nan")


def _compute_week_number_per_employee(df: pd.DataFrame) -> pd.Series:
    """
    1, 2, 3, … per employee_id by session time; only rows with employee_id, audit_id,
    and a present transcription.
    """
    out = pd.Series(pd.NA, index=df.index, dtype="Int64")
    mask = (
        df["employee_id"].map(_id_present)
        & df["audit_id"].map(_id_present)
        & df[COL_TRANSCRIPTION].map(_transcription_present)
    )
    if not mask.any():
        return out

    sub = df.loc[mask].copy()
    sub["_ts"] = sub.apply(_session_ts_utc, axis=1)
    sub = sub.sort_values(
        by=["employee_id", "_ts", "session_id"],
        kind="mergesort",
        na_position="last",
    )
    sub["_n"] = sub.groupby("employee_id", sort=False).cumcount() + 1
    out.loc[sub.index] = sub["_n"]
    return out


def _compute_audit_week(df: pd.DataFrame) -> pd.Series:
    """Calendar week index (1-based) since the first session week within each audit_id."""
    mon = df.apply(_session_ts_utc, axis=1).map(_monday_week_start_utc)
    out = pd.Series(pd.NA, index=df.index, dtype="Int64")

    for aid, grp in df.groupby("audit_id", sort=False):
        if not _id_present(aid):
            continue
        idx = grp.index
        m = mon.loc[idx]
        if m.isna().all():
            continue
        anchor = m.min()
        if pd.isna(anchor):
            continue
        weeks_elapsed = (m - anchor) // pd.Timedelta(days=7)
        out.loc[idx] = weeks_elapsed.astype("int64") + 1
    return out


def _is_missing_int(val: object) -> bool:
    if val is None or (isinstance(val, float) and pd.isna(val)):
        return True
    if isinstance(val, str) and not str(val).strip():
        return True
    try:
        return int(float(val)) < 1
    except (TypeError, ValueError):
        return True


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Fill week_number (per employee) and audit_week (per audit)."
    )
    parser.add_argument(
        "--input",
        type=Path,
        default=workbook_utils.DEFAULT_TRANSCRIPTS,
    )
    parser.add_argument(
        "--all-rows",
        action="store_true",
        help="Overwrite week_number and audit_week for every row that gets a computed value.",
    )
    args = parser.parse_args()
    path = args.input.expanduser().resolve()
    if not path.is_file():
        print(f"File not found: {path}", file=sys.stderr)
        return 1

    xl = pd.ExcelFile(path)
    sheet_name = workbook_utils.transcripts_sheet_name(list(xl.sheet_names))
    df = pd.read_excel(xl, sheet_name=sheet_name, header=HEADER_ROW - 1)

    for col in (
        "employee_id",
        "audit_id",
        COL_TRANSCRIPTION,
        "started_at",
        "ended_at",
        COL_WEEK_NUMBER,
        COL_AUDIT_WEEK,
    ):
        if col not in df.columns:
            print(f"Missing required column {col!r}.", file=sys.stderr)
            return 1

    week_num = _compute_week_number_per_employee(df)
    audit_week = _compute_audit_week(df)
    has_transcript = df[COL_TRANSCRIPTION].map(_transcription_present)

    if args.all_rows:
        needs_wn = has_transcript & week_num.notna()
        needs_wn_clear = ~has_transcript & ~df[COL_WEEK_NUMBER].map(_is_missing_int)
        needs_aw = audit_week.notna()
    else:
        needs_wn = df[COL_WEEK_NUMBER].map(_is_missing_int) & week_num.notna()
        needs_wn_clear = pd.Series(False, index=df.index)
        needs_aw = df[COL_AUDIT_WEEK].map(_is_missing_int) & audit_week.notna()

    if not needs_wn.any() and not needs_aw.any() and not needs_wn_clear.any():
        print(f"No updates needed in {path.name} ({len(df)} rows).")
        return 0

    wb = load_workbook(path)
    if sheet_name not in wb.sheetnames:
        print(f"Sheet {sheet_name!r} not found.", file=sys.stderr)
        return 1
    ws = wb[sheet_name]
    header = [ws.cell(HEADER_ROW, c).value for c in range(1, ws.max_column + 1)]
    try:
        col_wn = header.index(COL_WEEK_NUMBER) + 1
        col_aw = header.index(COL_AUDIT_WEEK) + 1
    except ValueError as e:
        print(f"Header column not found: {e}", file=sys.stderr)
        return 1

    n = len(df)
    n_wn = n_wn_clear = n_aw = 0
    for i in range(n):
        excel_row = DATA_START_ROW + i
        if needs_wn.iloc[i]:
            ws.cell(excel_row, col_wn).value = int(week_num.iloc[i])
            n_wn += 1
        elif needs_wn_clear.iloc[i]:
            ws.cell(excel_row, col_wn).value = None
            n_wn_clear += 1
        if needs_aw.iloc[i]:
            ws.cell(excel_row, col_aw).value = int(audit_week.iloc[i])
            n_aw += 1

    if n_wn == 0 and n_wn_clear == 0 and n_aw == 0:
        print("Nothing to write.")
        return 0

    workbook_utils.save_workbook_atomic(wb, path)
    parts = []
    if n_wn:
        parts.append(f"week_number={n_wn}")
    if n_wn_clear:
        parts.append(f"week_number cleared={n_wn_clear}")
    if n_aw:
        parts.append(f"audit_week={n_aw}")
    print(f"Updated {path.name}: {', '.join(parts)} row(s) (sheet {sheet_name!r}).")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
