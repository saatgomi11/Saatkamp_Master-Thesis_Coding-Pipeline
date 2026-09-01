"""Shared helpers for Gioia qualitative coding (Stages 1–2)."""

from __future__ import annotations

import json
import os
import re
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook

import workbook_utils

ROOT = workbook_utils.PROJECT_ROOT
DEFAULT_WORKBOOK = workbook_utils.DEFAULT_TRANSCRIPTS
QA_SHEET = "Qualitative Analysis"
QA_HEADER_ROW = 1
QA_DATA_START_ROW = QA_HEADER_ROW + 1

TRANSCRIPTS_HEADER_ROW = 2

PROMPT_STAGE1 = "Prompt_Gioia_Quotes"
PROMPT_STAGE2 = "Prompt_class_two"

STAGE2_JSON_OUT = ROOT / "outputs" / "processed" / "gioia_stage2_themes.json"
STAGE2_RAW_ERROR = ROOT / "outputs" / "processed" / "gioia_stage2_raw_error.txt"
STAGE2_CHECKPOINT = ROOT / "outputs" / "processed" / "gioia_stage2_checkpoint.json"

STAGE2_CONTINUATION_SUFFIX = """
---
CONTINUATION BATCH — COMPACT OUTPUT ONLY (batches 2+)
Do NOT return a "themes" array with first_order_codes lists.
Return ONLY valid JSON in this exact shape:

{
  "code_to_theme_map": {
    "<every code in first_order_codes_this_batch>": "<exact second_order_theme name>"
  },
  "new_themes": [
    {
      "second_order_theme": "<only if you created a new theme this batch>",
      "definition": "<one sentence>"
    }
  ]
}

"code_to_theme_map" must include every code in first_order_codes_this_batch exactly once.
"new_themes" may be [] if all codes map to existing_second_order_themes.
Reuse existing theme names character-for-character when they fit.
"""

TOPIC_SEGMENT_COLS: dict[str, str] = {
    "overall_experience": "text_overall_experience",
    "workload_sustainability": "text_workload_sustainability",
    "client_value": "text_client_value",
    "leadership_engagement": "text_leadership_engagement",
    "closing": "text_closing",
}

QA_COLUMNS = [
    "Session_ID",
    "Interview_Number",
    "Response_JSON",
    "Quote",
    "Topic",
    "Flaggin_Reason",
    "First_Code_Label",
    "Confidence_Score",
    "Second_Order",
    "My_Manual_Code",
    "Include_Flag",
]


def _load_dotenv() -> None:
    workbook_utils.load_project_dotenv()


def _present(raw: object) -> bool:
    if raw is None:
        return False
    try:
        if pd.isna(raw):
            return False
    except Exception:
        pass
    s = str(raw).strip().lower()
    return s not in ("", "nan", "none", "null")


def _load_prompt(name: str) -> str:
    p = workbook_utils.PROMPTS_DIR / name
    if not p.is_file():
        raise FileNotFoundError(f"Missing prompt file: {p}")
    return p.read_text(encoding="utf-8")


def _strip_json_fences(text: str) -> str:
    t = text.strip()
    if t.startswith("```"):
        t = re.sub(r"^```(?:json)?\s*", "", t, flags=re.IGNORECASE)
        t = re.sub(r"\s*```\s*$", "", t)
    return t.strip()


def _parse_model_json(text: str) -> dict[str, Any]:
    data = json.loads(_strip_json_fences(text))
    if not isinstance(data, dict):
        raise ValueError("Model JSON must be an object")
    return data


def _extract_braced_json_object(text: str, start: int) -> dict[str, Any] | None:
    if start < 0 or start >= len(text) or text[start] != "{":
        return None
    depth = 0
    in_str = False
    esc = False
    for i in range(start, len(text)):
        ch = text[i]
        if in_str:
            if esc:
                esc = False
            elif ch == "\\":
                esc = True
            elif ch == '"':
                in_str = False
            continue
        if ch == '"':
            in_str = True
        elif ch == "{":
            depth += 1
        elif ch == "}":
            depth -= 1
            if depth == 0:
                try:
                    obj = json.loads(text[start : i + 1])
                    return obj if isinstance(obj, dict) else None
                except json.JSONDecodeError:
                    return None
    return None


def _extract_json_object_for_key(text: str, key: str) -> dict[str, Any] | None:
    m = re.search(rf'"{re.escape(key)}"\s*:\s*\{{', text)
    if not m:
        return None
    return _extract_braced_json_object(text, m.end() - 1)


def _repair_stage2_json_text(text: str) -> str:
    """Fix common model mistake: themes array closed with }, instead of ]."""
    t = _strip_json_fences(text)
    t = re.sub(
        r'\n(\s*)\},\s*\n(\s*"new_themes"\s*:)',
        r"\n\1],\n\2",
        t,
        count=1,
    )
    t = re.sub(
        r'\n(\s*)\},\s*\n(\s*"code_to_theme_map"\s*:)',
        r"\n\1],\n\2",
        t,
        count=1,
    )
    return t


def _parse_stage2_json(text: str) -> dict[str, Any]:
    """
    Parse Stage 2 model output; repair or salvage code_to_theme_map if needed.
    """
    t = _strip_json_fences(text)
    for candidate in (t, _repair_stage2_json_text(t)):
        try:
            data = json.loads(candidate)
            if isinstance(data, dict):
                return data
        except json.JSONDecodeError:
            continue
    cmap = _extract_json_object_for_key(t, "code_to_theme_map")
    if cmap:
        new_themes: list[Any] = []
        m = re.search(r'"new_themes"\s*:\s*\[', t)
        if m:
            start = m.end() - 1
            depth = 0
            in_str = False
            esc = False
            for i in range(start, len(t)):
                ch = t[i]
                if in_str:
                    if esc:
                        esc = False
                    elif ch == "\\":
                        esc = True
                    elif ch == '"':
                        in_str = False
                    continue
                if ch == '"':
                    in_str = True
                elif ch == "[":
                    depth += 1
                elif ch == "]":
                    depth -= 1
                    if depth == 0:
                        try:
                            arr = json.loads(t[start : i + 1])
                            if isinstance(arr, list):
                                new_themes = arr
                        except json.JSONDecodeError:
                            pass
                        break
        return {"code_to_theme_map": cmap, "new_themes": new_themes, "_salvaged": True}
    raise json.JSONDecodeError("Could not parse or salvage Stage 2 JSON", t, 0)


def _user_excerpt_from_segment_json(raw: object) -> str:
    if not _present(raw):
        return ""
    try:
        data = json.loads(str(raw).strip())
    except json.JSONDecodeError:
        return str(raw).strip()
    if not isinstance(data, dict):
        return ""
    msgs = data.get("messages")
    if not isinstance(msgs, list):
        return ""
    parts: list[str] = []
    for m in msgs:
        if not isinstance(m, dict):
            continue
        if str(m.get("role", "")).lower() != "user":
            continue
        t = m.get("text")
        if t is None:
            continue
        chunk = str(t).strip()
        if chunk:
            parts.append(chunk)
    return "\n\n".join(parts).strip()


def build_topic_speech_payload(row: pd.Series) -> dict[str, str]:
    """Participant user speech per Gioia topic label."""
    out: dict[str, str] = {}
    for topic_key, col in TOPIC_SEGMENT_COLS.items():
        if col in row.index:
            out[topic_key] = _user_excerpt_from_segment_json(row.get(col))
    return out


def load_transcript_sessions(workbook: Path | None = None) -> tuple[pd.DataFrame, str]:
    path = workbook or DEFAULT_WORKBOOK
    xl = pd.ExcelFile(path)
    sheet = workbook_utils.transcripts_sheet_name(list(xl.sheet_names))
    df = pd.read_excel(xl, sheet_name=sheet, header=TRANSCRIPTS_HEADER_ROW - 1)
    df = df[df["session_id"].map(_present) & df["transcription"].map(_present)].copy()
    return df, sheet


def _normalize_quote_key(session_id: str, quote: str) -> tuple[str, str]:
    return str(session_id).strip(), re.sub(r"\s+", " ", str(quote or "").strip().lower())


def read_qualitative_sheet(workbook: Path | None = None) -> pd.DataFrame:
    path = workbook or DEFAULT_WORKBOOK
    xl = pd.ExcelFile(path)
    if QA_SHEET not in xl.sheet_names:
        return pd.DataFrame(columns=QA_COLUMNS)
    df = pd.read_excel(xl, sheet_name=QA_SHEET, header=QA_HEADER_ROW - 1)
    for col in QA_COLUMNS:
        if col not in df.columns:
            df[col] = None
    return df


def manual_override_index(df_qa: pd.DataFrame) -> dict[tuple[str, str], dict[str, Any]]:
    """Preserve human override columns keyed by (session_id, quote)."""
    out: dict[tuple[str, str], dict[str, Any]] = {}
    if df_qa.empty:
        return out
    for _, r in df_qa.iterrows():
        sid = r.get("Session_ID")
        quote = r.get("Quote")
        if not _present(sid):
            continue
        key = _normalize_quote_key(str(sid), str(quote or ""))
        out[key] = {
            "My_Manual_Code": r.get("My_Manual_Code"),
            "Include_Flag": r.get("Include_Flag"),
            "Second_Order": r.get("Second_Order"),
        }
    return out


def sessions_already_in_qa(df_qa: pd.DataFrame) -> set[str]:
    if df_qa.empty:
        return set()
    s = df_qa["Session_ID"].dropna().astype(str).str.strip()
    s = s[s != ""]
    has_json = df_qa["Response_JSON"].map(_present)
    return set(s[has_json].tolist())


def stage1_rows_from_response(
    session_id: str,
    interview_number: object,
    response_obj: dict[str, Any],
    response_json_text: str,
    overrides: dict[tuple[str, str], dict[str, Any]],
) -> list[dict[str, Any]]:
    quotes = response_obj.get("quotes") or []
    if not isinstance(quotes, list):
        quotes = []
    rows: list[dict[str, Any]] = []
    for q in quotes:
        if not isinstance(q, dict):
            continue
        quote_text = str(q.get("quote_text") or "").strip()
        if not quote_text:
            continue
        row = {
            "Session_ID": session_id,
            "Interview_Number": interview_number,
            "Response_JSON": response_json_text,
            "Quote": quote_text,
            "Topic": str(q.get("topic") or "").strip(),
            "Flaggin_Reason": str(q.get("flagging_reason") or "").strip(),
            "First_Code_Label": str(q.get("first_order_code") or "").strip(),
            "Confidence_Score": q.get("confidence"),
            "Second_Order": None,
            "My_Manual_Code": None,
            "Include_Flag": None,
        }
        key = _normalize_quote_key(session_id, quote_text)
        if key in overrides:
            for field in ("My_Manual_Code", "Include_Flag", "Second_Order"):
                val = overrides[key].get(field)
                if _present(val):
                    row[field] = val
        rows.append(row)
    return rows


def write_qualitative_sheet(df_qa: pd.DataFrame, workbook: Path | None = None) -> None:
    path = workbook or DEFAULT_WORKBOOK
    wb = load_workbook(path)
    if QA_SHEET not in wb.sheetnames:
        raise ValueError(f"Sheet {QA_SHEET!r} not found in {path.name}")
    ws = wb[QA_SHEET]

    # Ensure header row matches expected columns.
    for c, name in enumerate(QA_COLUMNS, start=1):
        ws.cell(QA_HEADER_ROW, c).value = name

    # Clear existing data rows.
    if ws.max_row >= QA_DATA_START_ROW:
        ws.delete_rows(QA_DATA_START_ROW, ws.max_row - QA_DATA_START_ROW + 1)

    out = df_qa.reindex(columns=QA_COLUMNS)
    for i, (_, r) in enumerate(out.iterrows()):
        excel_row = QA_DATA_START_ROW + i
        for c, col in enumerate(QA_COLUMNS, start=1):
            val = r.get(col)
            if pd.isna(val):
                val = None
            ws.cell(excel_row, c).value = val

    workbook_utils.save_workbook_atomic(wb, path)


def build_stage1_user_prompt(template: str, session_id: str, interview_number: object, speech: dict[str, str]) -> str:
    payload = {
        "session_id": session_id,
        "interview_number": int(interview_number) if _present(interview_number) else None,
        "participant_speech_by_topic": speech,
    }
    return (
        f"{template.strip()}\n\n"
        "---\n"
        "SESSION INPUT (JSON)\n"
        f"{json.dumps(payload, ensure_ascii=False, indent=2)}"
    )


def build_stage2_user_prompt(
    template: str,
    codes: list[str],
    *,
    existing_themes: list[dict[str, str]] | None = None,
    batch_index: int | None = None,
    batch_total: int | None = None,
) -> str:
    payload: dict[str, Any] = {
        "first_order_codes_this_batch": codes,
        "existing_second_order_themes": existing_themes or [],
    }
    if batch_index is not None and batch_total is not None:
        payload["batch"] = {"index": batch_index, "total": batch_total}
    parts = [
        template.strip(),
        "---\nSESSION INPUT (JSON)\n" + json.dumps(payload, ensure_ascii=False, indent=2),
    ]
    if existing_themes:
        parts.append(STAGE2_CONTINUATION_SUFFIX.strip())
    return "\n\n".join(parts)


def merge_theme_registry(
    registry: list[dict[str, str]],
    response: dict[str, Any],
) -> list[dict[str, str]]:
    """Accumulate unique second-order themes across sequential batches."""
    by_name: dict[str, dict[str, str]] = {
        str(t["second_order_theme"]).strip(): {
            "second_order_theme": str(t["second_order_theme"]).strip(),
            "definition": str(t.get("definition") or "").strip(),
        }
        for t in registry
        if _present(t.get("second_order_theme"))
    }
    sources: list[Any] = list(response.get("new_themes") or [])
    sources.extend(response.get("themes") or [])
    for th in sources:
        if not isinstance(th, dict):
            continue
        name = str(th.get("second_order_theme") or "").strip()
        if not name or name in by_name:
            continue
        by_name[name] = {
            "second_order_theme": name,
            "definition": str(th.get("definition") or "").strip(),
        }
    return sorted(by_name.values(), key=lambda x: x["second_order_theme"].lower())


def anthropic_client():
    _load_dotenv()
    try:
        import anthropic
    except ImportError as e:
        raise ImportError("Install: pip install anthropic python-dotenv") from e
    api_key = os.getenv("ANTHROPIC_API_KEY", "").strip()
    if not api_key:
        raise RuntimeError("Set ANTHROPIC_API_KEY in env/.env")
    model = os.getenv("ANTHROPIC_MODEL", "claude-sonnet-4-20250514").strip()
    max_tokens = int(os.getenv("ANTHROPIC_MAX_TOKENS", "16384"))
    sleep_s = float(os.getenv("ANTHROPIC_SLEEP_SEC", "0.35"))
    return anthropic.Anthropic(api_key=api_key), model, max_tokens, sleep_s


def call_anthropic_text(client: Any, model: str, max_tokens: int, user_prompt: str) -> tuple[str, str]:
    msg = client.messages.create(
        model=model,
        max_tokens=max_tokens,
        messages=[{"role": "user", "content": user_prompt}],
    )
    text_out = ""
    for block in msg.content:
        if hasattr(block, "text"):
            text_out += block.text
    stop = str(getattr(msg, "stop_reason", "") or "")
    return text_out, stop


def call_anthropic_json(
    client: Any,
    model: str,
    max_tokens: int,
    user_prompt: str,
    *,
    raw_error_path: Path | None = None,
    stage2: bool = False,
) -> dict[str, Any]:
    text_out, stop = call_anthropic_text(client, model, max_tokens, user_prompt)
    parser = _parse_stage2_json if stage2 else _parse_model_json
    try:
        data = parser(text_out)
        if data.get("_salvaged"):
            print("  Note: salvaged code_to_theme_map from partially malformed JSON.", file=__import__("sys").stderr)
        return data
    except json.JSONDecodeError as e:
        if raw_error_path is not None:
            raw_error_path.parent.mkdir(parents=True, exist_ok=True)
            raw_error_path.write_text(text_out, encoding="utf-8")
        hint = f" (stop_reason={stop!r}, chars={len(text_out)})" if stop else f" (chars={len(text_out)})"
        if stop == "max_tokens":
            hint += " — response likely truncated; use smaller --batch-size or raise ANTHROPIC_STAGE2_MAX_TOKENS."
        raise json.JSONDecodeError(e.msg + hint, e.doc, e.pos) from e


def save_stage2_checkpoint(
    codes: list[str],
    batch_size: int,
    last_batch_completed: int,
    merged_map: dict[str, str],
    theme_registry: list[dict[str, str]],
    batch_responses: list[dict[str, Any]],
) -> None:
    STAGE2_CHECKPOINT.parent.mkdir(parents=True, exist_ok=True)
    STAGE2_CHECKPOINT.write_text(
        json.dumps(
            {
                "codes": codes,
                "batch_size": batch_size,
                "last_batch_completed": last_batch_completed,
                "merged_map": merged_map,
                "theme_registry": theme_registry,
                "batch_responses": batch_responses,
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )


def load_stage2_checkpoint() -> dict[str, Any] | None:
    if not STAGE2_CHECKPOINT.is_file():
        return None
    try:
        return json.loads(STAGE2_CHECKPOINT.read_text(encoding="utf-8"))
    except Exception:
        return None


def extract_code_to_theme_map(data: dict[str, Any]) -> dict[str, str]:
    code_map: dict[str, str] = {}
    raw_map = data.get("code_to_theme_map")
    if isinstance(raw_map, dict):
        for k, v in raw_map.items():
            if _present(k) and _present(v):
                code_map[str(k).strip()] = str(v).strip()
    if not code_map:
        themes = data.get("themes") or []
        if isinstance(themes, list):
            for th in themes:
                if not isinstance(th, dict):
                    continue
                theme_name = str(th.get("second_order_theme") or "").strip()
                for code in th.get("first_order_codes") or []:
                    if _present(code) and theme_name:
                        code_map[str(code).strip()] = theme_name
    return code_map


def stage2_max_tokens() -> int:
    _load_dotenv()
    return int(os.getenv("ANTHROPIC_STAGE2_MAX_TOKENS", os.getenv("ANTHROPIC_MAX_TOKENS", "16384")))
