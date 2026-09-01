"""Core logic for manual disclosure-gap coding (reviewer 2; Excel I/O, scoring, session list)."""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook

import workbook_utils

ROOT = workbook_utils.PROJECT_ROOT
DEFAULT_WORKBOOK = workbook_utils.DEFAULT_TRANSCRIPTS
TOPICS_CONFIG = workbook_utils.CONFIG_DIR / "manual_gap_topics_reviewer2.json"
STATE_FILE = ROOT / ".manual_gap_coder_state_reviewer2.json"

HEADER_ROW = 2
DATA_START_ROW = HEADER_ROW + 1

LIKERT_MAP = {
    "disagree": 1,
    "tend_to_disagree": 2,
    "tend_to_agree": 3,
    "agree": 4,
}

SEVERITY_LABELS = {3: "Explicit", 2: "Hedged", 1: "Absent"}


def _present(raw: object) -> bool:
    if raw is None:
        return False
    try:
        if pd.isna(raw):
            return False
    except Exception:
        pass
    s = str(raw).strip().lower()
    return s not in ("", "nan", "none", "null")


def load_topics_config() -> list[dict[str, Any]]:
    data = json.loads(TOPICS_CONFIG.read_text(encoding="utf-8"))
    return list(data["topics"])


def parse_pre_likert(pre_json_raw: object, question_id: str) -> int | None:
    if not _present(pre_json_raw):
        return None
    try:
        d = json.loads(str(pre_json_raw))
    except Exception:
        return None
    if not isinstance(d, dict):
        return None
    item = d.get(question_id)
    if not isinstance(item, dict):
        return None
    val = str(item.get("value", "")).strip().lower()
    return LIKERT_MAP.get(val)


def compute_gap_segmentation(likert: int, transcript_severity: int) -> str:
    """Codebook matrix: product = likert x transcript severity (1-3)."""
    if likert <= 2:
        return "None"
    product = likert * transcript_severity
    if product >= 9:
        return "Strong"
    if product >= 5:
        return "Moderate"
    return "None"


def load_sessions(workbook: Path | None = None) -> tuple[pd.DataFrame, str]:
    path = workbook or DEFAULT_WORKBOOK
    xl = pd.ExcelFile(path)
    sheet = workbook_utils.transcripts_sheet_name(list(xl.sheet_names))
    df = pd.read_excel(xl, sheet_name=sheet, header=HEADER_ROW - 1)
    df["_excel_row"] = df.index + DATA_START_ROW
    df = df[df["session_id"].map(_present)].copy()
    df["_ts"] = pd.to_datetime(df["started_at"], utc=True, errors="coerce")
    df["_ts"] = df["_ts"].where(
        df["_ts"].notna(),
        pd.to_datetime(df["ended_at"], utc=True, errors="coerce"),
    )
    df = df.sort_values(["_ts", "session_id"], kind="mergesort", na_position="last")
    df = df.reset_index(drop=True)
    df["_session_num"] = df.index + 1
    return df, sheet


def _cell_value(raw: object) -> object:
    if not _present(raw):
        return None
    try:
        if pd.isna(raw):
            return None
    except Exception:
        pass
    return raw


def read_topic_values(row: pd.Series, topic: dict[str, Any]) -> dict[str, Any]:
    gap = _cell_value(row.get(topic["manual_gap"]))
    likert = _cell_value(row.get(topic["manual_likert"]))
    seg = _cell_value(row.get(topic["gap_segmentation"]))
    neg = _cell_value(row.get(topic["negative_issue"]))
    try:
        gap_n = int(float(gap)) if gap is not None else None
    except (TypeError, ValueError):
        gap_n = None
    try:
        likert_n = int(float(likert)) if likert is not None else None
    except (TypeError, ValueError):
        likert_n = None
    neg_s = str(neg).strip() if neg is not None else None
    if neg_s and neg_s.lower() in ("yes", "no"):
        neg_s = neg_s.capitalize()
    return {
        "manual_gap": gap_n,
        "manual_likert": likert_n,
        "gap_segmentation": str(seg).strip() if seg is not None else None,
        "negative_issue": neg_s,
    }


def find_resume_position(sessions: pd.DataFrame, topics: list[dict[str, Any]]) -> tuple[int, int]:
    for si in range(len(sessions)):
        row = sessions.iloc[si]
        for ti, topic in enumerate(topics):
            v = read_topic_values(row, topic)
            if v["manual_gap"] is None or v["negative_issue"] is None:
                return si, ti
    return max(0, len(sessions) - 1), 0


def load_state() -> dict[str, Any] | None:
    if not STATE_FILE.exists():
        return None
    try:
        return json.loads(STATE_FILE.read_text(encoding="utf-8"))
    except Exception:
        return None


def save_state(session_index: int, topic_index: int) -> None:
    STATE_FILE.write_text(
        json.dumps(
            {"session_index": session_index, "topic_index": topic_index},
            indent=2,
        ),
        encoding="utf-8",
    )


def clear_state() -> None:
    if STATE_FILE.exists():
        STATE_FILE.unlink()


def write_cells(
    excel_row: int,
    sheet_name: str,
    workbook: Path | None,
    updates: dict[str, object],
) -> None:
    path = workbook or DEFAULT_WORKBOOK
    wb = load_workbook(path)
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"Sheet {sheet_name!r} not found")
    ws = wb[sheet_name]
    header = [ws.cell(HEADER_ROW, c).value for c in range(1, ws.max_column + 1)]
    for col_name, value in updates.items():
        try:
            col_idx = header.index(col_name) + 1
        except ValueError as e:
            raise ValueError(f"Column not found in workbook: {col_name!r}") from e
        ws.cell(excel_row, col_idx).value = value
    workbook_utils.save_workbook_atomic(wb, path)


def save_severity(
    row: pd.Series,
    topic: dict[str, Any],
    severity: int,
    sheet_name: str,
    workbook: Path | None = None,
) -> dict[str, object]:
    likert = parse_pre_likert(row.get("pre-survey_response_json"), topic["pre_survey_question_id"])
    if likert is None:
        raise ValueError("Pre-survey Likert missing for this topic - fill pre-survey_response_json first.")
    seg = compute_gap_segmentation(likert, severity)
    updates = {
        topic["manual_gap"]: severity,
        topic["manual_likert"]: likert,
        topic["gap_segmentation"]: seg,
    }
    write_cells(int(row["_excel_row"]), sheet_name, workbook, updates)
    return {**updates, "likert": likert, "segmentation": seg}


def topic_text_to_chat_html(raw: object) -> str:
    """
    Render segmented topic JSON as a chat UI, or plain text if not JSON.
    Expected shape: {"name": "...", "messages": [{"role": "agent"|"user", "text": "..."}, ...]}
    """
    if not _present(raw):
        return '<div class="chat-empty"><em>(empty)</em></div>'

    text = str(raw).strip()
    try:
        data = json.loads(text)
    except Exception:
        return f'<div class="chat-plain">{_html_escape(text)}</div>'

    title = ""
    messages: list[dict[str, Any]] = []

    if isinstance(data, dict):
        title = str(data.get("name") or data.get("topic") or "").strip()
        msgs = data.get("messages")
        if isinstance(msgs, list):
            messages = [m for m in msgs if isinstance(m, dict)]
    elif isinstance(data, list):
        messages = [m for m in data if isinstance(m, dict)]

    if not messages:
        return f'<div class="chat-plain">{_html_escape(text)}</div>'

    parts: list[str] = ['<div class="chat-wrap">']
    if title:
        parts.append(f'<div class="chat-topic-title">{_html_escape(title)}</div>')
    for msg in messages:
        role = str(msg.get("role", "unknown")).strip().lower()
        body = str(msg.get("text", "") or "").strip()
        if not body and role not in ("agent", "user"):
            continue
        if not body:
            body = "..."
        label = "Agent" if role == "agent" else "Participant" if role == "user" else role.title()
        css = "chat-bubble-agent" if role == "agent" else "chat-bubble-user" if role == "user" else "chat-bubble-other"
        parts.append(
            f'<div class="chat-row {css}">'
            f'<div class="chat-role">{_html_escape(label)}</div>'
            f'<div class="chat-text">{_html_escape(body)}</div>'
            f"</div>"
        )
    parts.append("</div>")
    return "".join(parts)


def _html_escape(s: str) -> str:
    return (
        s.replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace('"', "&quot;")
        .replace("\n", "<br>")
    )


def save_negative_issue(
    row: pd.Series,
    topic: dict[str, Any],
    yes_no: str,
    sheet_name: str,
    workbook: Path | None = None,
) -> None:
    val = "Yes" if yes_no.lower() == "yes" else "No"
    write_cells(int(row["_excel_row"]), sheet_name, workbook, {topic["negative_issue"]: val})
