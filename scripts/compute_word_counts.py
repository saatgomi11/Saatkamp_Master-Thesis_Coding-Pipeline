#!/usr/bin/env python3
"""
Fill word-count columns from participant text and segmented topic JSON.

- word_count_total: word count of participant_text_clean (all user answers, no agent lines).
- word_count_overall, word_count_workload_sustainability, word_count_client_value,
  word_count_leadership_engagement, word_count_additional_topic, word_count_closing:
  from the matching text_* column JSON, count words in messages with role \"user\" only.

Skips rows with no transcription (same rule as other scripts). By default only writes a
count cell when it is still empty. Use --all-rows to overwrite all listed counts.

Close Excel before running.

Usage:
  python compute_word_counts.py
  python compute_word_counts.py --all-rows
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

import workbook_utils

HEADER_ROW = 2
DATA_START_ROW = HEADER_ROW + 1

COL_TRANSCRIPTION = "transcription"
COL_PARTICIPANT_CLEAN = "participant_text_clean"
COL_WORD_TOTAL = "word_count_total"

# Segmented JSON column -> word count column (see spreadsheet headers).
SEGMENT_TO_WORD: dict[str, str] = {
    "text_overall_experience": "word_count_overall",
    "text_workload_sustainability": "word_count_workload_sustainability",
    "text_client_value": "word_count_client_value",
    "text_leadership_engagement": "word_count_leadership_engagement",
    "text_additional_topic": "word_count_additional_topic",
    "text_closing": "word_count_closing",
}


def _transcription_present(raw: object) -> bool:
    if raw is None or (isinstance(raw, float) and pd.isna(raw)):
        return False
    s = str(raw).strip()
    return bool(s) and s.lower() not in ("null", "none", "nan")


def _is_missing_count(val: object) -> bool:
    if val is None or (isinstance(val, float) and pd.isna(val)):
        return True
    if isinstance(val, str):
        s = val.strip()
        if not s or s.lower() in ("nan", "none", "null"):
            return True
        try:
            return float(s) < 0
        except ValueError:
            return True
    try:
        return int(val) < 0  # type: ignore[arg-type]
    except (TypeError, ValueError):
        return True


def _count_words(text: str) -> int:
    if not text or not str(text).strip():
        return 0
    return len(str(text).split())


def _user_word_count_from_topic_json(raw: object) -> int | None:
    """Return None if there is no usable JSON object (skip writing)."""
    if raw is None or (isinstance(raw, float) and pd.isna(raw)):
        return None
    s = str(raw).strip()
    if not s or s.lower() in ("null", "none", "nan"):
        return None
    try:
        data = json.loads(s)
    except json.JSONDecodeError:
        return None
    if not isinstance(data, dict):
        return None
    msgs = data.get("messages")
    if msgs is None:
        msgs = []
    if not isinstance(msgs, list):
        return None
    parts: list[str] = []
    for m in msgs:
        if not isinstance(m, dict):
            continue
        if str(m.get("role", "")).lower() != "user":
            continue
        t = m.get("text")
        if t is None:
            continue
        parts.append(str(t))
    return _count_words(" ".join(parts))


def main() -> int:
    parser = argparse.ArgumentParser(description="Fill word count columns from clean text and topic JSON.")
    parser.add_argument(
        "--input",
        type=Path,
        default=workbook_utils.DEFAULT_TRANSCRIPTS,
        help="Path to Master-Supabase Transcripts.xlsx",
    )
    parser.add_argument(
        "--all-rows",
        action="store_true",
        help="Overwrite all target word-count cells when source data exists.",
    )
    args = parser.parse_args()
    path = args.input.expanduser().resolve()
    if not path.is_file():
        print(f"File not found: {path}", file=sys.stderr)
        return 1

    xl = pd.ExcelFile(path)
    sheet_name = workbook_utils.transcripts_sheet_name(list(xl.sheet_names))
    df = pd.read_excel(xl, sheet_name=sheet_name, header=HEADER_ROW - 1)

    required = [COL_TRANSCRIPTION, COL_PARTICIPANT_CLEAN, COL_WORD_TOTAL, *SEGMENT_TO_WORD.keys(), *SEGMENT_TO_WORD.values()]
    missing = [c for c in required if c not in df.columns]
    if missing:
        print(f"Missing columns: {missing}", file=sys.stderr)
        return 1

    wb = load_workbook(path)
    if sheet_name not in wb.sheetnames:
        print(f"Sheet {sheet_name!r} not found.", file=sys.stderr)
        return 1
    ws = wb[sheet_name]
    header = [ws.cell(HEADER_ROW, c).value for c in range(1, ws.max_column + 1)]

    def col_letter_idx(name: str) -> int:
        return header.index(name) + 1

    try:
        idx_total = col_letter_idx(COL_WORD_TOTAL)
        idx_map = {word_col: col_letter_idx(word_col) for word_col in SEGMENT_TO_WORD.values()}
    except ValueError as e:
        print(f"Header lookup failed: {e}", file=sys.stderr)
        return 1

    n = len(df)
    updates = 0

    for i in range(n):
        if not _transcription_present(df[COL_TRANSCRIPTION].iloc[i]):
            continue

        excel_row = DATA_START_ROW + i

        if args.all_rows or _is_missing_count(df[COL_WORD_TOTAL].iloc[i]):
            clean = df[COL_PARTICIPANT_CLEAN].iloc[i]
            if clean is None or (isinstance(clean, float) and pd.isna(clean)):
                wt = 0
            else:
                wt = _count_words(str(clean))
            ws.cell(excel_row, idx_total).value = int(wt)
            updates += 1

        for seg_col, word_col in SEGMENT_TO_WORD.items():
            if not args.all_rows and not _is_missing_count(df[word_col].iloc[i]):
                continue
            wc = _user_word_count_from_topic_json(df[seg_col].iloc[i])
            if wc is None:
                continue
            ws.cell(excel_row, idx_map[word_col]).value = int(wc)
            updates += 1

    if updates == 0:
        print(f"No word-count updates needed in {path.name} ({n} rows).")
        return 0

    workbook_utils.save_workbook_atomic(wb, path)
    print(f"Updated {path.name}: wrote {updates} word-count cell(s) on sheet {sheet_name!r}.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
