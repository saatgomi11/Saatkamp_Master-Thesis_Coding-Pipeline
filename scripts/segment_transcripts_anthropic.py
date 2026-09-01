#!/usr/bin/env python3
"""
Segment full interview JSON (transcription) into six topic columns via Anthropic.

Uses the prompt in ``Prompt_Segregation`` (placeholder ``<<INSERT TRANSCRIPT JSON HERE>>``).
Writes each topic's ``messages`` (and ``name``) as JSON text into:

  text_overall_experience, text_client_value, text_workload_sustainability,
  text_leadership_engagement, text_closing, text_additional_topic (maps ``other``)

By default processes a row when ``transcription`` is present and any of the six topic
cells is still empty. Use ``--all-rows`` to re-segment and overwrite when transcription exists.

Optional (no Anthropic): ``--refresh-negative-issue`` recomputes
``negative_issue_keyword_count`` and ``negative_issue_auto`` in Python from participant
text + existing sentiment columns (legacy helper; not part of the segregation prompt).

API key: copy ``env/.env.example`` to ``env/.env`` and set ANTHROPIC_API_KEY (see env/README).

Close Excel before running.

Usage:
  python segment_transcripts_anthropic.py
  python segment_transcripts_anthropic.py --limit 3
  python segment_transcripts_anthropic.py --dry-run
  python segment_transcripts_anthropic.py --all-rows
  python segment_transcripts_anthropic.py --refresh-negative-issue
"""

from __future__ import annotations

import argparse
import json
import os
import re
import sys
import time
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

import workbook_utils

try:
    import anthropic
except ImportError as e:  # pragma: no cover
    anthropic = None  # type: ignore

HEADER_ROW = 2
DATA_START_ROW = HEADER_ROW + 1

PROMPT_FILE = "Prompt_Segregation"
PLACEHOLDER = "<<INSERT TRANSCRIPT JSON HERE>>"

COL_TRANSCRIPTION = "transcription"

TOPIC_TO_COL = {
    "overall_experience": "text_overall_experience",
    "client_value": "text_client_value",
    "workload_sustainability": "text_workload_sustainability",
    "leadership_engagement": "text_leadership_engagement",
    "closing": "text_closing",
    "other": "text_additional_topic",
}

COL_PARTICIPANT_CLEAN = "participant_text_clean"
COL_SENTIMENT_NEG_AVG = "sentiment_negative_avg"
COL_SENTIMENT_NEG_MAX = "sentiment_negative_max"
COL_NEGATIVE_ISSUE_KW = "negative_issue_keyword_count"
COL_NEGATIVE_ISSUE_AUTO = "negative_issue_auto"

NEGATIVE_ISSUE_KEYWORDS: tuple[str, ...] = tuple(
    sorted(
        {
            "problem",
            "issue",
            "challenge",
            "difficult",
            "unclear",
            "inefficient",
            "inefficiency",
            "frustration",
            "frustrating",
            "pressure",
            "overload",
            "workload",
            "stress",
            "late",
            "overtime",
            "rework",
            "misaligned",
            "not aligned",
            "bottleneck",
            "blocker",
            "conflict",
            "concern",
            "ambiguity",
            "ambiguous",
            "delay",
            "delayed",
            "unsustainable",
            "unpredictable",
            "lack",
            "missing",
            "unclear priorities",
        },
        key=len,
        reverse=True,
    )
)

EXCEL_CELL_CHAR_LIMIT = 32700


def _load_dotenv() -> None:
    workbook_utils.load_project_dotenv()


def _transcription_present(raw: object) -> bool:
    if raw is None or (isinstance(raw, float) and pd.isna(raw)):
        return False
    s = str(raw).strip()
    return bool(s) and s.lower() not in ("null", "none", "nan")


def _is_empty_cell(val: object) -> bool:
    if val is None or (isinstance(val, float) and pd.isna(val)):
        return True
    s = str(val).strip()
    if not s or s.lower() in ("nan", "none", "null"):
        return True
    return False


def _is_missing_count(val: object) -> bool:
    if val is None or (isinstance(val, float) and pd.isna(val)):
        return True
    s = str(val).strip()
    if not s or s.lower() in ("nan", "none", "null"):
        return True
    try:
        float(s)
        return False
    except ValueError:
        return True


def _load_prompt_template() -> str:
    p = workbook_utils.PROMPTS_DIR / PROMPT_FILE
    if not p.is_file():
        raise FileNotFoundError(f"Missing prompt file: {p}")
    return p.read_text(encoding="utf-8")


def _build_user_prompt(template: str, transcript_json: str) -> str:
    if PLACEHOLDER not in template:
        raise ValueError(f"Prompt must contain placeholder {PLACEHOLDER!r}")
    return template.replace(PLACEHOLDER, transcript_json.strip(), 1)


def _parse_model_json(text: str) -> dict:
    t = text.strip()
    if t.startswith("```"):
        t = re.sub(r"^```(?:json)?\s*", "", t, flags=re.IGNORECASE)
        t = re.sub(r"\s*```\s*$", "", t)
    return json.loads(t)


def _topic_cell_value(topic_obj: object) -> str:
    if topic_obj is None:
        return json.dumps({"name": "", "messages": []}, ensure_ascii=False)
    if isinstance(topic_obj, dict):
        out = json.dumps(
            {"name": topic_obj.get("name", ""), "messages": topic_obj.get("messages", [])},
            ensure_ascii=False,
        )
    else:
        out = json.dumps({"name": "", "messages": []}, ensure_ascii=False)
    if len(out) > EXCEL_CELL_CHAR_LIMIT:
        out = out[: EXCEL_CELL_CHAR_LIMIT - 20] + "\n…[truncated]"
    return out


def _char_is_word_char(ch: str) -> bool:
    return bool(ch) and (ch.isalnum() or ch == "_")


def _negative_issue_keyword_count_from_text(text: str) -> int:
    """Greedy longest-first non-overlapping matches at word boundaries."""
    if not text or not str(text).strip():
        return 0
    s = re.sub(r"\s+", " ", str(text).strip().lower())
    if not s:
        return 0
    used = [False] * len(s)
    total = 0
    i = 0
    while i < len(s):
        if used[i]:
            i += 1
            continue
        matched_len = 0
        for kw in NEGATIVE_ISSUE_KEYWORDS:
            L = len(kw)
            if i + L > len(s):
                continue
            if s[i : i + L] != kw:
                continue
            before_ok = i == 0 or not _char_is_word_char(s[i - 1])
            after_ok = i + L >= len(s) or not _char_is_word_char(s[i + L])
            if before_ok and after_ok:
                matched_len = L
                break
        if matched_len:
            for j in range(i, i + matched_len):
                used[j] = True
            total += 1
            i += matched_len
        else:
            i += 1
    return total


def _safe_float(val: object) -> float:
    if val is None or (isinstance(val, float) and pd.isna(val)):
        return 0.0
    try:
        return float(val)
    except (TypeError, ValueError):
        return 0.0


def _negative_issue_auto_value(kw_count: int, neg_avg: object, neg_max: object) -> int:
    if kw_count >= 1:
        return 1
    if _safe_float(neg_avg) >= 0.50 or _safe_float(neg_max) >= 0.70:
        return 1
    return 0


def _participant_text_for_negative_issue(row: pd.Series) -> str:
    clean = row.get(COL_PARTICIPANT_CLEAN)
    if clean is not None and not (isinstance(clean, float) and pd.isna(clean)):
        t = str(clean).strip()
        if t and t.lower() not in ("null", "none", "nan"):
            return t
    raw = row.get(COL_TRANSCRIPTION)
    if not _transcription_present(raw):
        return ""
    try:
        data = json.loads(str(raw).strip())
    except json.JSONDecodeError:
        return ""
    msgs = data.get("messages")
    if not isinstance(msgs, list):
        return ""
    parts: list[str] = []
    for m in msgs:
        if not isinstance(m, dict):
            continue
        if str(m.get("role", "")).lower() != "user":
            continue
        t = m.get("text")
        if t is None:
            continue
        chunk = str(t).strip()
        if chunk:
            parts.append(chunk)
    return "\n\n".join(parts).strip()


def _write_negative_issue_cells(
    ws: object,
    excel_row: int,
    row: pd.Series,
    col_idx: dict[str, int],
) -> None:
    text = _participant_text_for_negative_issue(row)
    kw = _negative_issue_keyword_count_from_text(text)
    auto = _negative_issue_auto_value(kw, row.get(COL_SENTIMENT_NEG_AVG), row.get(COL_SENTIMENT_NEG_MAX))
    ws.cell(excel_row, col_idx[COL_NEGATIVE_ISSUE_KW]).value = int(kw)
    ws.cell(excel_row, col_idx[COL_NEGATIVE_ISSUE_AUTO]).value = int(auto)


def refresh_negative_issue_only(
    path: Path,
    *,
    dry_run: bool = False,
    limit: int = 0,
) -> int:
    """
    Recompute negative_issue_keyword_count and negative_issue_auto for every row with
    transcription (no Anthropic). Uses existing sentiment_negative_avg /
    sentiment_negative_max columns when present.
    """
    path = path.expanduser().resolve()
    if not path.is_file():
        print(f"File not found: {path}", file=sys.stderr)
        return 1

    xl = pd.ExcelFile(path)
    sheet_name = workbook_utils.transcripts_sheet_name(list(xl.sheet_names))
    df = pd.read_excel(xl, sheet_name=sheet_name, header=HEADER_ROW - 1)

    need = (
        COL_TRANSCRIPTION,
        COL_PARTICIPANT_CLEAN,
        COL_SENTIMENT_NEG_AVG,
        COL_SENTIMENT_NEG_MAX,
        COL_NEGATIVE_ISSUE_KW,
        COL_NEGATIVE_ISSUE_AUTO,
    )
    missing = [c for c in need if c not in df.columns]
    if missing:
        print(f"Missing columns: {missing}", file=sys.stderr)
        return 1

    indices = [i for i in range(len(df)) if _transcription_present(df[COL_TRANSCRIPTION].iloc[i])]
    if limit and limit > 0:
        indices = indices[:limit]

    if not indices:
        print("No rows with transcription to refresh.")
        return 0

    if dry_run:
        print(f"Dry run: would refresh negative_issue_* on {len(indices)} row(s).")
        return 0

    wb = load_workbook(path)
    if sheet_name not in wb.sheetnames:
        print(f"Sheet {sheet_name!r} not found.", file=sys.stderr)
        return 1
    ws = wb[sheet_name]
    header = [ws.cell(HEADER_ROW, c).value for c in range(1, ws.max_column + 1)]
    try:
        col_idx = {
            COL_NEGATIVE_ISSUE_KW: header.index(COL_NEGATIVE_ISSUE_KW) + 1,
            COL_NEGATIVE_ISSUE_AUTO: header.index(COL_NEGATIVE_ISSUE_AUTO) + 1,
        }
    except ValueError as e:
        print(f"Missing header column: {e}", file=sys.stderr)
        return 1

    for i in indices:
        excel_row = DATA_START_ROW + i
        _write_negative_issue_cells(ws, excel_row, df.iloc[i], col_idx)

    workbook_utils.save_workbook_atomic(wb, path)
    print(f"Refreshed negative_issue_* on {len(indices)} row(s) in {path.name} (sheet {sheet_name!r}).")
    return 0


def main() -> int:
    _load_dotenv()

    parser = argparse.ArgumentParser(description="Segment transcripts with Anthropic API.")
    parser.add_argument(
        "--input",
        type=Path,
        default=workbook_utils.DEFAULT_TRANSCRIPTS,
        help="Path to Master-Supabase Transcripts.xlsx",
    )
    parser.add_argument(
        "--all-rows",
        action="store_true",
        help="Re-segment and overwrite topic columns when transcription exists.",
    )
    parser.add_argument(
        "--limit",
        type=int,
        default=0,
        help="Process at most N matching rows (0 = no limit).",
    )
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="Print rows that would run without calling the API.",
    )
    parser.add_argument(
        "--refresh-negative-issue",
        action="store_true",
        help="Only recompute negative_issue_keyword_count and negative_issue_auto from text "
        "+ sentiment columns (no Anthropic).",
    )
    args = parser.parse_args()

    path = args.input.expanduser().resolve()

    if args.refresh_negative_issue:
        if args.all_rows:
            print("--all-rows is ignored with --refresh-negative-issue.", file=sys.stderr)
        return refresh_negative_issue_only(path, dry_run=args.dry_run, limit=args.limit)

    if not path.is_file():
        print(f"File not found: {path}", file=sys.stderr)
        return 1

    api_key = os.getenv("ANTHROPIC_API_KEY", "").strip()

    model = os.getenv("ANTHROPIC_MODEL", "claude-sonnet-4-20250514").strip()
    max_tokens = int(os.getenv("ANTHROPIC_MAX_TOKENS", "16384"))
    sleep_s = float(os.getenv("ANTHROPIC_SLEEP_SEC", "0.35"))

    template = _load_prompt_template()

    xl = pd.ExcelFile(path)
    sheet_name = workbook_utils.transcripts_sheet_name(list(xl.sheet_names))
    df = pd.read_excel(xl, sheet_name=sheet_name, header=HEADER_ROW - 1)

    if COL_TRANSCRIPTION not in df.columns:
        print(f"Missing column {COL_TRANSCRIPTION!r}.", file=sys.stderr)
        return 1
    for col in TOPIC_TO_COL.values():
        if col not in df.columns:
            print(f"Missing column {col!r}.", file=sys.stderr)
            return 1

    topic_cols = list(TOPIC_TO_COL.values())

    def row_topics_complete(i: int) -> bool:
        return not any(_is_empty_cell(df[c].iloc[i]) for c in topic_cols)

    def row_needs_anthropic(i: int) -> bool:
        if not _transcription_present(df[COL_TRANSCRIPTION].iloc[i]):
            return False
        if args.all_rows:
            return True
        return not row_topics_complete(i)

    indices = [i for i in range(len(df)) if row_needs_anthropic(i)]
    if args.limit and args.limit > 0:
        indices = indices[: args.limit]

    if not indices:
        print("No rows to process (missing transcription, or all six topic columns already filled).")
        return 0

    if args.dry_run:
        pairs = [f"df_idx={i}, excel_row={DATA_START_ROW + i}" for i in indices[:20]]
        tail = "" if len(indices) <= 20 else f" … (+{len(indices) - 20} more)"
        print(f"Dry run: would process {len(indices)} row(s). {', '.join(pairs)}{tail}")
        return 0

    if not api_key:
        print("Set ANTHROPIC_API_KEY in env/.env (see env/.env.example).", file=sys.stderr)
        return 1
    if anthropic is None:
        print("Install dependencies: pip install anthropic python-dotenv", file=sys.stderr)
        return 1

    wb = load_workbook(path)
    if sheet_name not in wb.sheetnames:
        print(f"Sheet {sheet_name!r} not found.", file=sys.stderr)
        return 1
    ws = wb[sheet_name]
    header = [ws.cell(HEADER_ROW, c).value for c in range(1, ws.max_column + 1)]
    try:
        col_idx = {col: header.index(col) + 1 for col in topic_cols}
    except ValueError as e:
        print(f"Missing header column: {e}", file=sys.stderr)
        return 1

    client = anthropic.Anthropic(api_key=api_key)
    processed = 0

    for i in indices:
        raw_tr = df[COL_TRANSCRIPTION].iloc[i]
        transcript_json = str(raw_tr).strip()
        user_prompt = _build_user_prompt(template, transcript_json)

        msg = client.messages.create(
            model=model,
            max_tokens=max_tokens,
            messages=[{"role": "user", "content": user_prompt}],
        )
        text_out = ""
        for block in msg.content:
            if hasattr(block, "text"):
                text_out += block.text
        data = _parse_model_json(text_out)

        excel_row = DATA_START_ROW + i
        for topic_key, col_name in TOPIC_TO_COL.items():
            cell_val = _topic_cell_value(data.get(topic_key))
            ws.cell(excel_row, col_idx[col_name]).value = cell_val

        workbook_utils.save_workbook_atomic(wb, path)
        processed += 1
        if sleep_s > 0:
            time.sleep(sleep_s)

    print(f"Done. Segmented {processed} row(s) in {path.name} (sheet {sheet_name!r}).")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
