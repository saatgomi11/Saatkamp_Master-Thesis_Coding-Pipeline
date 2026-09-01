#!/usr/bin/env python3
"""
Fill post_feedback_topic from the post-survey JSON (post-survey_response_json).

The post survey stores several keyed answers; the open comment is stored as a plain string.
This script takes the *last* non-empty string value when iterating the JSON object in key
order (matches schemas where the open question is the final field). Likert entries are ints
or dicts, so they are skipped.

Only writes when post_feedback_topic is empty unless --all-rows is passed.

Close Excel before running.

Usage:
  python extract_post_feedback_topic.py
  python extract_post_feedback_topic.py --all-rows
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

import workbook_utils

HEADER_ROW = 2
DATA_START_ROW = HEADER_ROW + 1

COL_JSON = "post-survey_response_json"
COL_TOPIC = "post_feedback_topic"


def _is_empty_cell(val: object) -> bool:
    if val is None or (isinstance(val, float) and pd.isna(val)):
        return True
    s = str(val).strip()
    if not s or s.lower() in ("nan", "none", "null"):
        return True
    return False


def _last_string_value_in_object_order(data: dict) -> str | None:
    """Last non-empty string among values() in insertion order."""
    out: str | None = None
    for v in data.values():
        if isinstance(v, str):
            t = v.strip()
            if t:
                out = t
    return out


def _topic_from_post_json(raw: object) -> str | None:
    if raw is None or (isinstance(raw, float) and pd.isna(raw)):
        return None
    s = str(raw).strip()
    if not s or s.lower() in ("null", "none", "nan"):
        return None
    try:
        data = json.loads(s)
    except json.JSONDecodeError:
        return None
    if not isinstance(data, dict):
        return None
    return _last_string_value_in_object_order(data)


def main() -> int:
    parser = argparse.ArgumentParser(description="Fill post_feedback_topic from post-survey JSON.")
    parser.add_argument(
        "--input",
        type=Path,
        default=workbook_utils.DEFAULT_TRANSCRIPTS,
        help="Path to Master-Supabase Transcripts.xlsx",
    )
    parser.add_argument(
        "--all-rows",
        action="store_true",
        help="Overwrite post_feedback_topic whenever post JSON yields a string.",
    )
    args = parser.parse_args()
    path = args.input.expanduser().resolve()
    if not path.is_file():
        print(f"File not found: {path}", file=sys.stderr)
        return 1

    xl = pd.ExcelFile(path)
    sheet_name = workbook_utils.transcripts_sheet_name(list(xl.sheet_names))
    df = pd.read_excel(xl, sheet_name=sheet_name, header=HEADER_ROW - 1)

    for col in (COL_JSON, COL_TOPIC):
        if col not in df.columns:
            print(f"Missing column {col!r}.", file=sys.stderr)
            return 1

    computed = df[COL_JSON].map(_topic_from_post_json)
    has_text = computed.map(lambda x: isinstance(x, str) and bool(x.strip()))

    if args.all_rows:
        needs = has_text
    else:
        needs = df[COL_TOPIC].map(_is_empty_cell) & has_text

    if not needs.any():
        print(f"No post_feedback_topic updates needed in {path.name} ({len(df)} rows).")
        return 0

    wb = load_workbook(path)
    if sheet_name not in wb.sheetnames:
        print(f"Sheet {sheet_name!r} not found.", file=sys.stderr)
        return 1
    ws = wb[sheet_name]
    header = [ws.cell(HEADER_ROW, c).value for c in range(1, ws.max_column + 1)]
    try:
        col_topic = header.index(COL_TOPIC) + 1
    except ValueError:
        print(f"Column {COL_TOPIC!r} not found in header row.", file=sys.stderr)
        return 1

    n_written = 0
    for i in range(len(df)):
        if not needs.iloc[i]:
            continue
        text = computed.iloc[i]
        if not text:
            continue
        ws.cell(DATA_START_ROW + i, col_topic).value = text
        n_written += 1

    workbook_utils.save_workbook_atomic(wb, path)
    print(f"Updated {path.name}: wrote post_feedback_topic for {n_written} row(s).")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
