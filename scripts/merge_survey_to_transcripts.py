#!/usr/bin/env python3
"""
Merge survey results into the transcripts workbook by session_id.

For each transcript row, looks up matching rows in Master-Supabase Survey Results.xlsx
(survey_phase pre / post), copies response_json and submitted_at into:
  pre-survey_response_json, pre-survey_submitted_at
  post-survey_response_json, post-survey_submitted_at

By default only writes cells that are still empty (new sessions). If a session has only a
pre or only a post survey in the source table, the other side is left unchanged.

If multiple survey rows exist for the same session_id + phase, the latest submitted_at wins.

pre-survey_submitted_at and post-survey_submitted_at are written as plain text (original string
from the survey export when available, otherwise ISO-8601). That avoids Excel/openpyxl
timezone issues; you do not need to change column formats in Excel.

Close Excel before running.

Transcript sheet name is resolved like ``process_transcripts.py`` (see ``workbook_utils``).
Saves atomically (temp file then replace) to reduce half-written workbooks on errors.

Usage:
  python merge_survey_to_transcripts.py
  python merge_survey_to_transcripts.py --all-rows
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

import workbook_utils

HEADER_ROW = 2
DATA_START_ROW = HEADER_ROW + 1

COL_PRE_JSON = "pre-survey_response_json"
COL_PRE_SUB = "pre-survey_submitted_at"
COL_POST_JSON = "post-survey_response_json"
COL_POST_SUB = "post-survey_submitted_at"


def _is_empty_cell(val: object) -> bool:
    if val is None:
        return True
    if isinstance(val, float) and pd.isna(val):
        return True
    if isinstance(val, pd.Timestamp):
        return False
    s = str(val).strip()
    if not s or s.lower() in ("nan", "none", "null"):
        return True
    return False


def _submitted_as_string(row: pd.Series) -> str | None:
    """Store submitted_at as text so openpyxl never sees tz-aware datetimes."""
    raw = row.get("submitted_at")
    if raw is not None and not (isinstance(raw, float) and pd.isna(raw)):
        if isinstance(raw, pd.Timestamp):
            if pd.isna(raw):
                pass
            else:
                return raw.isoformat()
        s = str(raw).strip()
        if s and s.lower() not in ("nan", "none", "null", "nat"):
            return s
    ts = row.get("_submitted")
    if ts is None or (isinstance(ts, float) and pd.isna(ts)):
        return None
    tsn = pd.Timestamp(ts)
    if pd.isna(tsn):
        return None
    return tsn.isoformat()


def _normalize_response_json(val: object) -> str | None:
    if val is None or (isinstance(val, float) and pd.isna(val)):
        return None
    if isinstance(val, str):
        t = val.strip()
        return t if t and t.lower() not in ("null", "none", "nan") else None
    return str(val)


def _build_survey_lookup(survey_df: pd.DataFrame) -> tuple[dict[str, tuple[str, str | None]], dict[str, tuple[str, str | None]]]:
    """Returns (pre_by_session, post_by_session) -> session_id -> (response_json_str, submitted_at_str)."""
    required = {"session_id", "survey_phase", "response_json", "submitted_at"}
    missing = required - set(survey_df.columns)
    if missing:
        raise ValueError(f"Survey sheet missing columns: {sorted(missing)}")

    s = survey_df.copy()
    s["_phase"] = s["survey_phase"].astype(str).str.strip().str.lower()
    s = s[s["_phase"].isin(("pre", "post"))]
    s["_submitted"] = pd.to_datetime(s["submitted_at"], utc=True, errors="coerce")
    s = s.sort_values("_submitted", ascending=False, na_position="last")
    s = s.drop_duplicates(subset=["session_id", "_phase"], keep="first")

    pre: dict[str, tuple[str, str | None]] = {}
    post: dict[str, tuple[str, str | None]] = {}
    for _, row in s.iterrows():
        sid = row["session_id"]
        if pd.isna(sid) or sid is None:
            continue
        sid_str = str(sid).strip()
        if not sid_str:
            continue
        rj = _normalize_response_json(row["response_json"])
        if rj is None:
            continue
        sub = _submitted_as_string(row)
        if row["_phase"] == "pre":
            pre[sid_str] = (rj, sub)
        else:
            post[sid_str] = (rj, sub)
    return pre, post


def main() -> int:
    parser = argparse.ArgumentParser(description="Merge survey Excel into transcripts Excel.")
    parser.add_argument(
        "--transcripts",
        type=Path,
        default=workbook_utils.DEFAULT_TRANSCRIPTS,
        help="Path to Master-Supabase Transcripts.xlsx",
    )
    parser.add_argument(
        "--survey",
        type=Path,
        default=workbook_utils.DEFAULT_SURVEY,
        help="Path to Master-Supabase Survey Results.xlsx",
    )
    parser.add_argument(
        "--all-rows",
        action="store_true",
        help="Overwrite all four survey columns from the survey file when a match exists.",
    )
    args = parser.parse_args()

    t_path = args.transcripts.expanduser().resolve()
    s_path = args.survey.expanduser().resolve()
    if not t_path.is_file():
        print(f"Transcripts file not found: {t_path}", file=sys.stderr)
        return 1
    if not s_path.is_file():
        print(f"Survey file not found: {s_path}", file=sys.stderr)
        return 1

    survey = pd.read_excel(s_path, header=1)
    pre_map, post_map = _build_survey_lookup(survey)

    txl = pd.ExcelFile(t_path)
    sheet_name = workbook_utils.transcripts_sheet_name(list(txl.sheet_names))
    df = pd.read_excel(txl, sheet_name=sheet_name, header=HEADER_ROW - 1)
    if "session_id" not in df.columns:
        print("Transcripts sheet missing session_id.", file=sys.stderr)
        return 1
    for c in (COL_PRE_JSON, COL_PRE_SUB, COL_POST_JSON, COL_POST_SUB):
        if c not in df.columns:
            print(f"Transcripts sheet missing column {c!r}.", file=sys.stderr)
            return 1

    wb = load_workbook(t_path)
    if sheet_name not in wb.sheetnames:
        print(f"Sheet {sheet_name!r} not found. Available: {wb.sheetnames}", file=sys.stderr)
        return 1
    ws = wb[sheet_name]

    header = [ws.cell(HEADER_ROW, c).value for c in range(1, ws.max_column + 1)]
    try:
        col_idx = {
            COL_PRE_JSON: header.index(COL_PRE_JSON) + 1,
            COL_PRE_SUB: header.index(COL_PRE_SUB) + 1,
            COL_POST_JSON: header.index(COL_POST_JSON) + 1,
            COL_POST_SUB: header.index(COL_POST_SUB) + 1,
        }
    except ValueError as e:
        print(f"Could not find survey columns in header row: {e}", file=sys.stderr)
        return 1

    n = len(df)
    n_pre_j = n_pre_s = n_post_j = n_post_s = 0

    for i in range(n):
        sid = df["session_id"].iloc[i]
        if pd.isna(sid):
            continue
        sid_str = str(sid).strip()
        if not sid_str:
            continue

        pre = pre_map.get(sid_str)
        post = post_map.get(sid_str)

        excel_row = DATA_START_ROW + i

        if pre is not None:
            rj, sub = pre
            wj = args.all_rows or _is_empty_cell(df[COL_PRE_JSON].iloc[i])
            wsj = args.all_rows or _is_empty_cell(df[COL_PRE_SUB].iloc[i])
            if wj:
                ws.cell(excel_row, col_idx[COL_PRE_JSON]).value = rj
                n_pre_j += 1
            if wsj:
                ws.cell(excel_row, col_idx[COL_PRE_SUB]).value = sub
                n_pre_s += 1

        if post is not None:
            rj, sub = post
            wj = args.all_rows or _is_empty_cell(df[COL_POST_JSON].iloc[i])
            wsj = args.all_rows or _is_empty_cell(df[COL_POST_SUB].iloc[i])
            if wj:
                ws.cell(excel_row, col_idx[COL_POST_JSON]).value = rj
                n_post_j += 1
            if wsj:
                ws.cell(excel_row, col_idx[COL_POST_SUB]).value = sub
                n_post_s += 1

    total_writes = n_pre_j + n_pre_s + n_post_j + n_post_s
    if total_writes == 0:
        print(f"No survey fields to update in {t_path.name} ({n} transcript rows). Nothing to write.")
        return 0

    workbook_utils.save_workbook_atomic(wb, t_path)
    print(
        f"Updated {t_path.name}: pre JSON={n_pre_j}, pre submitted_at={n_pre_s}, "
        f"post JSON={n_post_j}, post submitted_at={n_post_s} (cell writes). Rows: {n}."
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
