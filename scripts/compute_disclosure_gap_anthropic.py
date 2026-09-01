#!/usr/bin/env python3
"""
Compare pre-survey Likert answers to interview speech per topic (disclosure gap) via Anthropic.

Uses ``Prompt_Disclosure_Gap`` and ``config/pre_survey_likert_map.json`` (UUID → topic mapping).
Writes:

  Disclosure gap              — Yes / No (any topic with a gap)
  gap_overall_expierience / reason_o_e  — overall experience
  gap_workload_sust / reason_w_s
  gap_client_value / reason_c_v
  gap_leadership_eng / reason_l_e

Requires: transcription, pre-survey_response_json; uses text_* topic JSON when present.
By default only rows with empty ``Disclosure gap `` cell (or missing per-topic gap). Use --all-rows
to overwrite.

Close Excel before running.

Usage:
  python compute_disclosure_gap_anthropic.py
  python compute_disclosure_gap_anthropic.py --dry-run
  python compute_disclosure_gap_anthropic.py --limit 3
  python compute_disclosure_gap_anthropic.py --all-rows
"""

from __future__ import annotations

import argparse
import json
import os
import re
import sys
import time
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

import workbook_utils

try:
    import anthropic
except ImportError:  # pragma: no cover
    anthropic = None  # type: ignore

HEADER_ROW = 2
DATA_START_ROW = HEADER_ROW + 1

PROMPT_FILE = "Prompt_Disclosure_Gap"
MAP_FILE = "pre_survey_likert_map.json"

PLACEHOLDER_PRE = "<<INSERT PRE_SURVEY BY TOPIC JSON HERE>>"
PLACEHOLDER_TRANSCRIPT = "<<INSERT TRANSCRIPT JSON HERE>>"
PLACEHOLDER_SEGMENTS = "<<INSERT TOPIC SEGMENTS JSON HERE>>"

COL_TRANSCRIPTION = "transcription"
COL_PRE_JSON = "pre-survey_response_json"

# Session-level (note trailing space in spreadsheet header)
COL_DISCLOSURE_GAP_ANY = "Disclosure gap "

TOPIC_KEYS = (
    "overall_experience",
    "client_value",
    "workload_sustainability",
    "leadership_engagement",
)

# Model topic_key -> (gap column, reason column)
TOPIC_TO_EXCEL: dict[str, tuple[str, str]] = {
    "overall_experience": ("gap_overall_expierience", "reason_o_e"),
    "workload_sustainability": ("gap_workload_sust", "reason_w_s"),
    "client_value": ("gap_client_value", "reason_c_v"),
    "leadership_engagement": ("gap_leadership_eng", "reason_l_e"),
}

# Segmented columns for 1:1 survey ↔ interview topic
SEGMENT_COLS = {
    "overall_experience": "text_overall_experience",
    "client_value": "text_client_value",
    "workload_sustainability": "text_workload_sustainability",
    "leadership_engagement": "text_leadership_engagement",
}

VALID_GAP_TYPES = frozenset({"none", "mild", "moderate", "strong"})


def _load_dotenv() -> None:
    workbook_utils.load_project_dotenv()


def _load_likert_map() -> dict:
    p = workbook_utils.CONFIG_DIR / MAP_FILE
    if not p.is_file():
        raise FileNotFoundError(f"Missing mapping file: {p}")
    return json.loads(p.read_text(encoding="utf-8"))


def _load_prompt_template() -> str:
    p = workbook_utils.PROMPTS_DIR / PROMPT_FILE
    if not p.is_file():
        raise FileNotFoundError(f"Missing prompt file: {p}")
    return p.read_text(encoding="utf-8")


def _transcription_present(raw: object) -> bool:
    if raw is None or (isinstance(raw, float) and pd.isna(raw)):
        return False
    s = str(raw).strip()
    return bool(s) and s.lower() not in ("null", "none", "nan")


def _pre_survey_present(raw: object) -> bool:
    if raw is None or (isinstance(raw, float) and pd.isna(raw)):
        return False
    s = str(raw).strip()
    return bool(s) and s.lower() not in ("null", "none", "nan", "{}")


def _is_empty_cell(val: object) -> bool:
    if val is None or (isinstance(val, float) and pd.isna(val)):
        return True
    s = str(val).strip()
    return not s or s.lower() in ("nan", "none", "null")


def _likert_valence(value: str, likert_map: dict) -> str:
    v = value.lower().strip().replace(" ", "_")
    if v in likert_map.get("likert_positive_values", []):
        return "positive"
    if v in likert_map.get("likert_negative_values", []):
        return "negative"
    label = v.replace("_", " ")
    if "disagree" in label or label == "disagree":
        return "negative"
    if "agree" in label:
        return "positive"
    return "positive" if "agree" in v else "negative"


def _pre_survey_by_topic(raw: object, likert_map: dict) -> dict:
    """Build topic_key -> {question_id, question_text, likert_label, likert_value, likert_valence}."""
    out: dict[str, dict] = {}
    try:
        data = json.loads(str(raw).strip())
    except json.JSONDecodeError:
        return out
    if not isinstance(data, dict):
        return out
    questions = likert_map.get("questions") or {}
    for qid, meta in questions.items():
        if not isinstance(meta, dict):
            continue
        topic_key = meta.get("topic_key")
        if not topic_key or topic_key not in TOPIC_KEYS:
            continue
        entry = data.get(qid)
        if not isinstance(entry, dict):
            continue
        label = str(entry.get("label") or entry.get("value") or "").strip()
        value = str(entry.get("value") or label).strip()
        out[topic_key] = {
            "question_id": qid,
            "question_text": meta.get("question_text", ""),
            "likert_label": label,
            "likert_value": value,
            "likert_valence": _likert_valence(value, likert_map),
        }
    return out


def _user_excerpt_from_segment_json(raw: object) -> str | None:
    if raw is None or (isinstance(raw, float) and pd.isna(raw)):
        return None
    s = str(raw).strip()
    if not s or s.lower() in ("null", "none", "nan"):
        return None
    try:
        data = json.loads(s)
    except json.JSONDecodeError:
        return None
    if not isinstance(data, dict):
        return None
    msgs = data.get("messages")
    if not isinstance(msgs, list):
        return None
    parts: list[str] = []
    for m in msgs:
        if not isinstance(m, dict):
            continue
        if str(m.get("role", "")).lower() != "user":
            continue
        t = m.get("text")
        if t is None:
            continue
        chunk = str(t).strip()
        if chunk:
            parts.append(chunk)
    text = "\n\n".join(parts).strip()
    return text if text else None


def _topic_segments_payload(row: pd.Series) -> dict:
    """User-only excerpts per topic from segmented interview columns."""
    payload: dict[str, dict] = {}
    for topic_key, col in SEGMENT_COLS.items():
        if col not in row.index:
            continue
        excerpt = _user_excerpt_from_segment_json(row.get(col))
        payload[topic_key] = {
            "segment_column": col,
            "user_messages_excerpt": excerpt or "",
        }
    return payload


def _build_user_prompt(
    template: str,
    transcript_json: str,
    pre_by_topic: dict,
    segments: dict,
) -> str:
    for ph, content in (
        (PLACEHOLDER_PRE, json.dumps(pre_by_topic, ensure_ascii=False, indent=2)),
        (PLACEHOLDER_TRANSCRIPT, transcript_json.strip()),
        (PLACEHOLDER_SEGMENTS, json.dumps(segments, ensure_ascii=False, indent=2)),
    ):
        if ph not in template:
            raise ValueError(f"Prompt must contain placeholder {ph!r}")
        template = template.replace(ph, content, 1)
    return template


def _parse_model_json(text: str) -> dict:
    t = text.strip()
    if t.startswith("```"):
        t = re.sub(r"^```(?:json)?\s*", "", t, flags=re.IGNORECASE)
        t = re.sub(r"\s*```\s*$", "", t)
    return json.loads(t)


def _normalize_gap_type(raw: object) -> str:
    s = str(raw or "none").strip().lower()
    return s if s in VALID_GAP_TYPES else "none"


def _normalize_topic_block(raw: object) -> dict:
    if not isinstance(raw, dict):
        raw = {}
    return {
        "gap_type": _normalize_gap_type(raw.get("gap_type")),
        "reason": str(raw.get("reason") or "").strip(),
        "disclosure_gap": bool(raw.get("disclosure_gap")),
    }


def _disclosure_gap_any_label(data: dict) -> str:
    any_gap = data.get("disclosure_gap_any")
    if any_gap is None:
        for tk in TOPIC_KEYS:
            block = data.get(tk)
            if isinstance(block, dict) and block.get("disclosure_gap"):
                any_gap = True
                break
        else:
            any_gap = False
    return "Yes" if any_gap else "No"


def _row_needs_run(row: pd.Series, all_rows: bool) -> bool:
    if not _transcription_present(row.get(COL_TRANSCRIPTION)):
        return False
    if not _pre_survey_present(row.get(COL_PRE_JSON)):
        return False
    if all_rows:
        return True
    if _is_empty_cell(row.get(COL_DISCLOSURE_GAP_ANY)):
        return True
    for gap_col, _ in TOPIC_TO_EXCEL.values():
        if gap_col in row.index and _is_empty_cell(row.get(gap_col)):
            return True
    return False


def main() -> int:
    _load_dotenv()

    parser = argparse.ArgumentParser(description="Disclosure gap coding via Anthropic.")
    parser.add_argument(
        "--input",
        type=Path,
        default=workbook_utils.DEFAULT_TRANSCRIPTS,
    )
    parser.add_argument("--all-rows", action="store_true")
    parser.add_argument("--limit", type=int, default=0)
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()

    path = args.input.expanduser().resolve()
    if not path.is_file():
        print(f"File not found: {path}", file=sys.stderr)
        return 1

    likert_map = _load_likert_map()
    template = _load_prompt_template()

    xl = pd.ExcelFile(path)
    sheet_name = workbook_utils.transcripts_sheet_name(list(xl.sheet_names))
    df = pd.read_excel(xl, sheet_name=sheet_name, header=HEADER_ROW - 1)

    required = [COL_TRANSCRIPTION, COL_PRE_JSON, COL_DISCLOSURE_GAP_ANY, *SEGMENT_COLS.values()]
    for col in required:
        if col not in df.columns:
            print(f"Missing column {col!r}.", file=sys.stderr)
            return 1
    for gap_col, reason_col in TOPIC_TO_EXCEL.values():
        if gap_col not in df.columns or reason_col not in df.columns:
            print(f"Missing gap/reason columns: {gap_col!r}, {reason_col!r}", file=sys.stderr)
            return 1

    indices = [i for i in range(len(df)) if _row_needs_run(df.iloc[i], args.all_rows)]
    if args.limit and args.limit > 0:
        indices = indices[: args.limit]

    if not indices:
        print("No rows to process (need transcription + pre-survey + empty disclosure columns).")
        return 0

    if args.dry_run:
        pairs = [f"df_idx={i}, excel_row={DATA_START_ROW + i}" for i in indices[:20]]
        tail = "" if len(indices) <= 20 else f" … (+{len(indices) - 20} more)"
        print(f"Dry run: would process {len(indices)} row(s). {', '.join(pairs)}{tail}")
        return 0

    api_key = os.getenv("ANTHROPIC_API_KEY", "").strip()
    if not api_key:
        print("Set ANTHROPIC_API_KEY in env/.env (see env/.env.example).", file=sys.stderr)
        return 1
    if anthropic is None:
        print("Install dependencies: pip install anthropic python-dotenv", file=sys.stderr)
        return 1

    model = os.getenv("ANTHROPIC_MODEL", "claude-sonnet-4-20250514").strip()
    max_tokens = int(os.getenv("ANTHROPIC_MAX_TOKENS", "16384"))
    sleep_s = float(os.getenv("ANTHROPIC_SLEEP_SEC", "0.35"))

    wb = load_workbook(path)
    ws = wb[sheet_name]
    header = [ws.cell(HEADER_ROW, c).value for c in range(1, ws.max_column + 1)]

    def col_idx(name: str) -> int:
        return header.index(name) + 1

    idx: dict[str, int] = {COL_DISCLOSURE_GAP_ANY: col_idx(COL_DISCLOSURE_GAP_ANY)}
    for gap_col, reason_col in TOPIC_TO_EXCEL.values():
        idx[gap_col] = col_idx(gap_col)
        idx[reason_col] = col_idx(reason_col)

    client = anthropic.Anthropic(api_key=api_key)
    processed = 0

    for i in indices:
        row = df.iloc[i]
        pre_by_topic = _pre_survey_by_topic(row[COL_PRE_JSON], likert_map)
        if not pre_by_topic:
            print(f"Row {DATA_START_ROW + i}: no mapped pre-survey topics; skip.", file=sys.stderr)
            continue

        transcript_json = str(row[COL_TRANSCRIPTION]).strip()
        segments = _topic_segments_payload(row)
        user_prompt = _build_user_prompt(template, transcript_json, pre_by_topic, segments)

        msg = client.messages.create(
            model=model,
            max_tokens=max_tokens,
            messages=[{"role": "user", "content": user_prompt}],
        )
        text_out = ""
        for block in msg.content:
            if hasattr(block, "text"):
                text_out += block.text
        data = _parse_model_json(text_out)

        excel_row = DATA_START_ROW + i
        ws.cell(excel_row, idx[COL_DISCLOSURE_GAP_ANY]).value = _disclosure_gap_any_label(data)

        for topic_key, (gap_col, reason_col) in TOPIC_TO_EXCEL.items():
            raw_block = data.get(topic_key)
            # Older prompt versions used psychological_safety for this column
            if topic_key == "overall_experience" and raw_block is None:
                raw_block = data.get("psychological_safety")
            block = _normalize_topic_block(raw_block)
            ws.cell(excel_row, idx[gap_col]).value = block["gap_type"]
            ws.cell(excel_row, idx[reason_col]).value = block["reason"]

        workbook_utils.save_workbook_atomic(wb, path)
        processed += 1
        if sleep_s > 0:
            time.sleep(sleep_s)

    print(f"Done. Disclosure gap coded for {processed} row(s) in {path.name} (sheet {sheet_name!r}).")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
